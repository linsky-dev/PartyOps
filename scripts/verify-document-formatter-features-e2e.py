#!/usr/bin/env python3
"""在真实 WPS/Word 源码宿主上验证六类功能及配置隔离。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from uuid import uuid4

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

BEIJING = timezone(timedelta(hours=8))


def platform_name() -> str:
    if os.name == "nt":
        return "windows"
    return "macos" if sys.platform == "darwin" else "linux"


def architecture_name(target_platform: str) -> str:
    machine = platform.machine().lower()
    if target_platform == "windows":
        return "x86" if machine in {"x86", "i386", "i686"} or sys.maxsize <= 2**32 else "x64"
    if target_platform == "macos":
        return "arm64" if machine in {"arm64", "aarch64"} else "x86_64"
    return "arm64" if machine in {"arm64", "aarch64"} else "amd64"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def snapshot(paths: list[Path]) -> dict[str, str | None]:
    return {path.name: sha256(path) if path.is_file() else None for path in paths}


def docx_text(path: Path) -> str:
    document = Document(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(blocks)


def assert_docx(path: Path) -> None:
    if not path.is_file() or path.stat().st_size <= 0:
        raise AssertionError(f"DOCX 不存在或为空：{path}")
    with zipfile.ZipFile(path) as package:
        if "word/document.xml" not in package.namelist():
            raise AssertionError(f"DOCX 缺少 document.xml：{path}")
        etree.fromstring(package.read("word/document.xml"))


def create_word_fixture(path: Path) -> None:
    document = Document()
    title = document.add_paragraph("关于开展基层治理专项工作的通知")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("测试党发〔2026〕12号")
    document.add_paragraph("同志们：")
    document.add_paragraph("旧称发布于2026年。")
    document.add_paragraph("第12条")
    formatted = document.add_paragraph()
    run = formatted.add_run("格式文字")
    run.font.name = "宋体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(16)
    document.add_paragraph("这是用于转换和套红的正文内容。")
    document.save(path)


def create_pdf_fixture(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((72, 96), "PartyOps PDF 2026", fontsize=16)
    page.insert_text((72, 126), "Local source engine verification", fontsize=12)
    document.save(path)
    document.close()


def one_docx(outputs: tuple[Any, ...]) -> Path:
    matches = [item.path for item in outputs if item.path.suffix.lower() == ".docx"]
    if len(matches) != 1:
        raise AssertionError(f"预期一个 DOCX，实际为：{[item.path.name for item in outputs]}")
    assert_docx(matches[0])
    return matches[0]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--host", type=Path, required=True)
    parser.add_argument(
        "--host-identity",
        type=Path,
        help="仅在跨环境代理测试时指定实际受测宿主；正式目标机测试应省略。",
    )
    parser.add_argument("--workspace", type=Path, required=True)
    parser.add_argument("--evidence", type=Path, required=True)
    parser.add_argument("--platform", choices=("windows", "linux", "macos"))
    parser.add_argument("--architecture")
    args = parser.parse_args()

    root = args.root.resolve()
    host = args.host.resolve()
    host_identity = (args.host_identity or args.host).resolve()
    workspace = args.workspace.resolve()
    evidence = args.evidence.resolve()
    target_platform = args.platform or platform_name()
    target_architecture = args.architecture or architecture_name(target_platform)
    if not host.is_file() or not host_identity.is_file():
        raise FileNotFoundError(f"源码宿主或身份文件不存在：{host} / {host_identity}")
    if workspace == workspace.parent or workspace == root:
        raise ValueError("E2E 工作目录不能是磁盘根目录或 PartyOps 仓库根目录。")
    workspace.mkdir(parents=True, exist_ok=True)
    # 每次执行使用独立目录，避免 WPS 的自动重命名策略或上一次中断残留污染断言。
    run_workspace = workspace / f"run-{uuid4().hex}"
    run_workspace.mkdir(parents=True, exist_ok=False)
    backend = root / "backend"
    sys.path.insert(0, str(backend))
    from app.official_format_features import execute_feature

    os.environ["PARTYOPS_DOCUMENT_FORMATTER_HOST"] = str(host)
    fixture = run_workspace / "input.docx"
    pdf_fixture = run_workspace / "input.pdf"
    create_word_fixture(fixture)
    create_pdf_fixture(pdf_fixture)
    input_hashes = {
        path.name: sha256(path)
        for path in (fixture, pdf_fixture)
    }

    if target_platform == "windows":
        local_app_data = Path(os.environ["LOCALAPPDATA"]) / "DocumentRepository"
    elif target_platform == "macos":
        local_app_data = Path.home() / "Library" / "Application Support" / "DocumentRepository"
    else:
        local_app_data = Path(
            os.environ.get("XDG_CONFIG_HOME", str(Path.home() / ".config"))
        ) / "DocumentRepository"
    config_paths = [
        local_app_data / "Templates.xml",
        local_app_data / "ReplacePlans.xml",
        local_app_data / "RedHeaderTemplates.xml",
        local_app_data / "RenameRules.xml",
        local_app_data / "ConvertOptions.xml",
    ]
    results: list[dict[str, Any]] = []

    def run_case(
        name: str,
        feature_id: str,
        source: Path,
        options: dict[str, Any],
    ) -> tuple[Any, ...]:
        before = snapshot(config_paths)
        source_hash = sha256(source)
        events: list[dict[str, Any]] = []
        case_workspace = run_workspace / name
        result = execute_feature(
            feature_id,
            source,
            case_workspace,
            options,
            progress=lambda percent, message: events.append(
                {"percent": percent, "message": message}
            ),
            cancelled=lambda: False,
        )
        outputs = result.outputs
        after = snapshot(config_paths)
        if before != after:
            raise AssertionError(f"{name} 未恢复原排版设置：before={before}, after={after}")
        if sha256(source) != source_hash:
            raise AssertionError(f"{name} 修改了源文件")
        results.append(
            {
                "case": name,
                "feature_id": feature_id,
                "status": "passed",
                "host": "partyops-source-host+wps",
                "outputs": [
                    {
                        "name": item.path.name,
                        "bytes": item.path.stat().st_size,
                        "sha256": sha256(item.path),
                    }
                    for item in outputs
                ],
                "last_progress": events[-1] if events else None,
                "configuration_restored": True,
            }
        )
        return outputs

    golden_input = (
        root
        / "backend"
        / "tests"
        / "fixtures"
        / "document-formatter-source"
        / "input-manual-break.docx"
    )
    format_output = one_docx(
        run_case(
            "format",
            "format",
            golden_input,
            {"compatibility_mode": "auto", "template": "GB/T 9704-2012"},
        )
    )
    if len(Document(str(format_output)).paragraphs) != 8:
        raise AssertionError("一键排版没有按源码把手动换行拆为 8 个段落")

    replace_output = one_docx(
        run_case(
            "replace",
            "replace",
            fixture,
            {
                "compatibility_mode": "auto",
                "plan_name": "E2E 四模式",
                "rules": [
                    {"mode": "text", "find": "旧称", "replace": "新称"},
                    {"mode": "regex", "find": "2026年", "replace": "二〇二六年"},
                    {"mode": "wildcard", "find": "第[0-9]{1,}条", "replace": "条款"},
                    {
                        "mode": "format",
                        "find": "宋体",
                        "font_name": "黑体",
                        "font_size": 18,
                        "alignment": "center",
                    },
                ],
            },
        )
    )
    replaced = docx_text(replace_output)
    for expected in ("新称", "二〇二六年", "条款", "格式文字"):
        if expected not in replaced:
            raise AssertionError(f"一键替换缺少结果：{expected}")
    if "旧称" in replaced or "2026年" in replaced or "第12条" in replaced:
        raise AssertionError("一键替换仍残留原文字")
    with zipfile.ZipFile(replace_output) as package:
        replace_xml = package.read("word/document.xml").decode("utf-8")
    # WPS Linux 将同一“黑体”字体以英文内部名 SimHei 序列化；两者是
    # 平台别名，字号和可见字体槽仍必须由原源码写入。
    if not ({"黑体", "SimHei"} & set(replace_xml.split('"'))) or 'w:sz w:val="36"' not in replace_xml:
        raise AssertionError("格式替换未通过原源码写入黑体 18 磅")

    redheader_output = one_docx(
        run_case(
            "redheader",
            "redheader",
            fixture,
            {
                "compatibility_mode": "auto",
                "document_type": "down",
                "copy_number": "1",
                "security": "秘密★10年",
                "urgency": "加急",
                "agency": "中共测试市委文件",
                "document_number": "测试党发〔2026〕12号",
                "signatory": "",
                "imprint": "中共测试市委办公室",
            },
        )
    )
    redheader_text = docx_text(redheader_output)
    for expected in ("中共测试市委文件", "测试党发〔2026〕12号", "中共测试市委办公室"):
        if expected not in redheader_text:
            raise AssertionError(f"一键套红缺少内容：{expected}")

    rename_output = one_docx(
        run_case(
            "rename",
            "rename",
            fixture,
            {
                "compatibility_mode": "auto",
                "parts": ["title", "custom"],
                "custom_text": "定稿",
                "separator": "-",
                "rotation_words": "",
            },
        )
    )
    if "定稿" not in rename_output.stem or "-" not in rename_output.stem:
        raise AssertionError(f"一键命名未应用组合规则：{rename_output.name}")

    convert_expectations = [
        ("convert-docx", "docx", "pages", ".docx"),
        ("convert-pdf", "pdf", "pages", ".pdf"),
        ("convert-txt", "txt", "pages", ".txt"),
        ("convert-png-pages", "png", "pages", ".png"),
        ("convert-jpg-long", "jpg", "long", ".jpg"),
    ]
    for name, target, image_mode, extension in convert_expectations:
        outputs = run_case(
            name,
            "convert",
            fixture,
            {
                "compatibility_mode": "auto",
                "target_format": target,
                "image_mode": image_mode,
                "page_selection": "all",
                "dpi": 144,
                "same_name_policy": "auto-rename",
            },
        )
        matching = [item.path for item in outputs if item.path.suffix.lower() == extension]
        if not matching:
            raise AssertionError(f"一键转换 {target}/{image_mode} 未生成 {extension}")
        if extension == ".docx":
            assert_docx(matching[0])
        elif extension == ".pdf" and not matching[0].read_bytes().startswith(b"%PDF"):
            raise AssertionError("转换 PDF 文件头无效")
        elif extension == ".txt" and "正文内容" not in matching[0].read_text(encoding="utf-8"):
            raise AssertionError("转换 TXT 缺少正文")
        elif extension in {".png", ".jpg"}:
            pixmap = fitz.Pixmap(str(matching[0]))
            if pixmap.width <= 0 or pixmap.height <= 0:
                raise AssertionError("转换图片尺寸无效")

    pdf_to_word_output = one_docx(
        run_case(
            "pdf-to-word",
            "pdf-to-word",
            pdf_fixture,
            {"compatibility_mode": "auto", "normalize_punctuation": True},
        )
    )
    if "PartyOps" not in docx_text(pdf_to_word_output):
        raise AssertionError("PDF 转 Word 未还原可编辑文本")

    for source in (fixture, pdf_fixture):
        if sha256(source) != input_hashes[source.name]:
            raise AssertionError(f"E2E 输入文件被修改：{source.name}")

    evidence.parent.mkdir(parents=True, exist_ok=True)
    record = {
        "schema": 2,
        "status": "passed",
        "verified_at": datetime.now(BEIJING).replace(microsecond=0).isoformat(),
        "timezone": "Asia/Shanghai",
        "platform": target_platform,
        "architecture": target_architecture,
        "provider": "wps",
        "host_sha256": sha256(host_identity),
        "host": str(host),
        "host_identity": str(host_identity),
        "case_count": len(results),
        "cases": results,
    }
    # Python 3.8 的 Path.write_text 尚不支持 newline 参数；Windows 7 构建链
    # 固定使用 Python 3.8，因此通过 Path.open 显式保持证据文件为 LF。
    with evidence.open("w", encoding="utf-8", newline="\n") as stream:
        stream.write(json.dumps(record, ensure_ascii=False, indent=2) + "\n")
    print(f"[FORMATTER_FEATURE_E2E_OK] {len(results)} cases -> {evidence}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
