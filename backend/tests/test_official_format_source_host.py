"""原排版源码无窗口宿主的本机协议与路由测试。"""

from __future__ import annotations

import json
from pathlib import Path

import fitz
import pytest
from docx import Document

from app import official_format_features, official_format_host
from app.official_format import OfficialFormatError


class _CompletedProcess:
    returncode = 0

    def poll(self) -> int:
        return 0


@pytest.mark.parametrize("mode", ["standard", "custom", "empty", "missing", "incomplete"])
def test_linux_host_discovers_wps_libraries_without_user_environment(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, mode: str,
) -> None:
    """普通用户环境为空时也应发现 WPS 自带 Qt，不能借用测试 LD 路径。"""
    monkeypatch.setattr(official_format_host.sys, "platform", "linux")
    monkeypatch.setattr(official_format_host.sys, "frozen", True, raising=False)
    monkeypatch.setenv("LD_LIBRARY_PATH", "/app/_internal")
    monkeypatch.delenv("LD_LIBRARY_PATH_ORIG", raising=False)
    monkeypatch.delenv("PARTYOPS_WPS_RPC_LIBRARY", raising=False)
    office = tmp_path / "WPS 中文"
    (office / "lib").mkdir(parents=True)
    rpc = office / "librpcwpsapi_wpsqt.so"
    if mode not in {"missing", "incomplete"}:
        rpc.touch()
    if mode != "incomplete":
        (office / "wps").touch()
    monkeypatch.setattr(official_format_host, "_LINUX_WPS_ROOTS", (office,))
    if mode == "custom":
        monkeypatch.setenv("PARTYOPS_WPS_RPC_LIBRARY", str(rpc))
        monkeypatch.setattr(official_format_host, "_LINUX_WPS_ROOTS", ())
    if mode != "empty":
        monkeypatch.setenv("LD_LIBRARY_PATH_ORIG", "/user/lib")
    result = official_format_host._source_host_environment()
    if mode in {"missing", "incomplete"}:
        assert result["LD_LIBRARY_PATH"] == "/user/lib"
    else:
        expected = [str(office), str(office / "lib")]
        if mode != "empty":
            expected.append("/user/lib")
        assert result["LD_LIBRARY_PATH"] == ":".join(expected)
    assert official_format_host.os.environ["LD_LIBRARY_PATH"] == "/app/_internal"


@pytest.mark.parametrize(
    ("runtime_platform", "os_name", "machine", "maxsize", "expected"),
    [
        ("darwin", "posix", "arm64", 2**63 - 1, ("partyops-document-formatter-host",)),
        ("linux", "posix", "x86_64", 2**63 - 1, ("partyops-document-formatter-host",)),
        (
            "win32",
            "nt",
            "x86",
            2**63 - 1,
            ("PartyOps.DocumentFormatter.Host-x86.exe", "PartyOps.DocumentFormatter.Host.exe"),
        ),
        (
            "win32",
            "nt",
            "AMD64",
            2**63 - 1,
            (
                "PartyOps.DocumentFormatter.Host-x64.exe",
                "PartyOps.DocumentFormatter.Host.exe",
                "PartyOps.DocumentFormatter.Host-x86.exe",
            ),
        ),
    ],
)
def test_source_host_names_are_selected_for_the_target_runtime(
    monkeypatch: pytest.MonkeyPatch,
    runtime_platform: str,
    os_name: str,
    machine: str,
    maxsize: int,
    expected: tuple[str, ...],
) -> None:
    monkeypatch.setattr(official_format_host.sys, "platform", runtime_platform)
    monkeypatch.setattr(official_format_host.os, "name", os_name)
    monkeypatch.setattr(official_format_host.platform, "machine", lambda: machine)
    monkeypatch.setattr(official_format_host.sys, "maxsize", maxsize)
    assert official_format_host._executable_names() == expected


def test_source_host_candidates_prioritize_explicit_packaged_host(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    packaged_host = tmp_path / "relocated" / "PartyOps.DocumentFormatter.Host.exe"
    monkeypatch.setenv("PARTYOPS_DOCUMENT_FORMATTER_HOST", str(packaged_host))
    monkeypatch.setattr(
        official_format_host,
        "_executable_names",
        lambda: ("PartyOps.DocumentFormatter.Host.exe",),
    )
    assert official_format_host.source_host_candidates()[0] == packaged_host.resolve()


@pytest.mark.parametrize(
    ("runtime_platform", "frozen", "original", "expected_library_path"),
    [
        ("linux", True, "/wps/lib", "/wps/lib"),
        ("linux", True, None, None),
        ("linux", True, "", None),
        ("linux", False, "/wps/lib", "/frozen/_internal:/wps/lib"),
        ("win32", True, "/wps/lib", "/frozen/_internal:/wps/lib"),
        ("darwin", True, "/wps/lib", "/frozen/_internal:/wps/lib"),
    ],
)
def test_source_host_protocol_accepts_only_private_real_outputs(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
    runtime_platform: str, frozen: bool, original: str | None,
    expected_library_path: str | None,
) -> None:
    monkeypatch.setattr(official_format_host.sys, "platform", runtime_platform)
    monkeypatch.setattr(official_format_host.sys, "frozen", frozen, raising=False)
    monkeypatch.setenv("LD_LIBRARY_PATH", "/frozen/_internal:/wps/lib")
    if original is None:
        monkeypatch.delenv("LD_LIBRARY_PATH_ORIG", raising=False)
    else:
        monkeypatch.setenv("LD_LIBRARY_PATH_ORIG", original)
    host = tmp_path / "host.exe"
    host.write_bytes(b"MZ")
    source = tmp_path / "source.docx"
    Document().save(source)
    workspace = tmp_path / "workspace"
    events: list[tuple[int, str]] = []
    child_environment: dict[str, str] = {}
    monkeypatch.setattr(official_format_host, "resolve_source_host", lambda: host)

    def fake_popen(command: list[str], **kwargs: object) -> _CompletedProcess:
        child_environment.update(kwargs["env"])
        request = Path(command[command.index("--request") + 1])
        response = Path(command[command.index("--response") + 1])
        progress = Path(command[command.index("--progress") + 1])
        payload = json.loads(request.read_text(encoding="utf-8"))
        output = Path(payload["output_directory"]) / "source_已排版.docx"
        Document().save(output)
        progress.write_text(
            json.dumps({"percent": 72, "message": "源码规则完成"}, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        response.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "success_count": 1,
                    "failure_count": 0,
                    "cancelled_count": 0,
                    "jobs": [
                        {
                            "success": True,
                            "cancelled": False,
                            "message": "排版完成",
                            "host_display_name": "WPS Office",
                            "output_paths": [str(output)],
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        return _CompletedProcess()

    monkeypatch.setattr(official_format_host.subprocess, "Popen", fake_popen)
    outputs = official_format_host.run_source_host(
        "format",
        source,
        workspace,
        {"compatibility_mode": "auto"},
        progress=lambda percent, message: events.append((percent, message)),
        cancelled=lambda: False,
    )
    assert len(outputs) == 1
    assert outputs[0].path.name == "source_已排版.docx"
    assert outputs[0].host_display_name == "WPS Office"
    assert events == [(72, "源码规则完成")]
    request_payload = json.loads(
        (workspace / ".source-host" / "request.json").read_text(encoding="utf-8")
    )
    assert request_payload["host_preference"] == "wps-preferred"
    assert request_payload["options"] == {"compatibility_mode": "auto"}
    assert child_environment["PARTYOPS_WPS_VTABLE_MAP"] == str(
        host.parent / "word-vtable-map.json"
    )
    assert child_environment.get("LD_LIBRARY_PATH") == expected_library_path
    assert official_format_host.os.environ["LD_LIBRARY_PATH"] == "/frozen/_internal:/wps/lib"


def test_source_host_missing_and_outside_output_are_hard_failures(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.docx"
    Document().save(source)
    monkeypatch.setattr(official_format_host, "resolve_source_host", lambda: None)
    with pytest.raises(OfficialFormatError) as missing:
        official_format_host.run_source_host(
            "format",
            source,
            tmp_path / "missing",
            {},
            progress=lambda *_args: None,
            cancelled=lambda: False,
        )
    assert missing.value.code == "SOURCE_FORMATTER_HOST_MISSING"

    host = tmp_path / "host.exe"
    host.write_bytes(b"MZ")
    outside = tmp_path / "outside.docx"
    Document().save(outside)
    monkeypatch.setattr(official_format_host, "resolve_source_host", lambda: host)

    def fake_popen(command: list[str], **_kwargs: object) -> _CompletedProcess:
        response = Path(command[command.index("--response") + 1])
        response.write_text(
            json.dumps(
                {
                    "jobs": [
                        {
                            "success": True,
                            "cancelled": False,
                            "message": "完成",
                            "output_paths": [str(outside)],
                        }
                    ]
                }
            ),
            encoding="utf-8",
        )
        return _CompletedProcess()

    monkeypatch.setattr(official_format_host.subprocess, "Popen", fake_popen)
    with pytest.raises(OfficialFormatError) as invalid:
        official_format_host.run_source_host(
            "format",
            source,
            tmp_path / "outside-check",
            {},
            progress=lambda *_args: None,
            cancelled=lambda: False,
        )
    assert invalid.value.code == "SOURCE_FORMATTER_OUTPUT_INVALID"


def test_source_host_progress_ignores_damaged_events_and_clamps_percent(
    tmp_path: Path,
) -> None:
    progress_path = tmp_path / "progress.jsonl"
    progress_path.write_bytes(
        b"\xff\n"
        b"{broken json}\n"
        b'{"percent":"not-a-number"}\n'
        b'{"percent":150}\n'
    )
    events: list[tuple[int, str]] = []

    consumed = official_format_host._read_progress(
        progress_path,
        0,
        lambda percent, message: events.append((percent, message)),
    )

    assert consumed == progress_path.stat().st_size
    assert events == [(100, "正在处理文档")]


@pytest.mark.parametrize(
    ("response_text", "expected_code"),
    [
        (None, "SOURCE_FORMATTER_NO_RESPONSE"),
        ("{broken json", "SOURCE_FORMATTER_RESPONSE_INVALID"),
        (
            json.dumps({"fatal": True, "message": "WPS 启动失败"}, ensure_ascii=False),
            "SOURCE_FORMATTER_FATAL",
        ),
        (json.dumps({"jobs": []}), "SOURCE_FORMATTER_FAILED"),
        (
            json.dumps(
                {
                    "jobs": [
                        {
                            "success": False,
                            "cancelled": True,
                            "message": "用户取消",
                        }
                    ]
                },
                ensure_ascii=False,
            ),
            "FORMAT_JOB_CANCELLED",
        ),
        (
            json.dumps(
                {
                    "jobs": [
                        {
                            "success": True,
                            "cancelled": False,
                            "output_paths": [],
                        }
                    ]
                }
            ),
            "SOURCE_FORMATTER_OUTPUT_MISSING",
        ),
    ],
)
def test_source_host_rejects_incomplete_or_failed_responses(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    response_text: str | None,
    expected_code: str,
) -> None:
    host = tmp_path / "host.exe"
    host.write_bytes(b"MZ")
    source = tmp_path / "source.docx"
    Document().save(source)
    monkeypatch.setattr(official_format_host, "resolve_source_host", lambda: host)

    def fake_popen(command: list[str], **_kwargs: object) -> _CompletedProcess:
        if response_text is not None:
            response = Path(command[command.index("--response") + 1])
            response.write_text(response_text, encoding="utf-8")
        return _CompletedProcess()

    monkeypatch.setattr(official_format_host.subprocess, "Popen", fake_popen)
    with pytest.raises(OfficialFormatError) as failure:
        official_format_host.run_source_host(
            "format",
            source,
            tmp_path / expected_code.lower(),
            {"compatibility_mode": "auto"},
            progress=lambda *_args: None,
            cancelled=lambda: False,
        )
    assert failure.value.code == expected_code


def test_source_host_rejects_unknown_compatibility_mode(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    host = tmp_path / "host.exe"
    host.write_bytes(b"MZ")
    source = tmp_path / "source.docx"
    Document().save(source)
    monkeypatch.setattr(official_format_host, "resolve_source_host", lambda: host)

    with pytest.raises(OfficialFormatError) as failure:
        official_format_host.run_source_host(
            "format",
            source,
            tmp_path / "invalid-mode",
            {"compatibility_mode": "libreoffice"},
            progress=lambda *_args: None,
            cancelled=lambda: False,
        )
    assert failure.value.code == "FORMAT_COMPATIBILITY_INVALID"


@pytest.mark.parametrize(
    ("feature_id", "suffix", "options"),
    [
        ("format", ".docx", {"compatibility_mode": "wps"}),
        ("replace", ".docx", {"rules": [{"mode": "text", "find": "旧", "replace": "新"}]}),
        ("redheader", ".docx", {"document_type": "down"}),
        ("rename", ".docx", {"parts": ["title"]}),
        ("convert", ".docx", {"target_format": "pdf"}),
        ("pdf-to-word", ".pdf", {}),
    ],
)
def test_all_production_features_route_to_source_host(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    feature_id: str,
    suffix: str,
    options: dict[str, object],
) -> None:
    monkeypatch.delenv("PARTYOPS_FORMATTER_TEST_LOCAL_ENGINE", raising=False)
    source = tmp_path / f"source{suffix}"
    output = tmp_path / "host-output.docx"
    if suffix == ".docx":
        Document().save(source)
    else:
        source.write_bytes(b"%PDF-1.4\n%%EOF")
    Document().save(output)
    calls: list[tuple[str, Path]] = []

    def fake_host(feature_id: str, source_path: Path, _workspace: Path, *_args: object, **_kwargs: object):
        calls.append((feature_id, source_path))
        return (
            official_format_host.SourceHostOutput(
                path=output,
                host_display_name="WPS Office",
                message="排版完成",
            ),
        )

    monkeypatch.setattr(official_format_host, "run_source_host", fake_host)
    result = official_format_features.execute_feature(
        feature_id,
        source,
        tmp_path / "job",
        options,
    )
    assert calls == [(feature_id, source)]
    assert result.outputs[0].path == output


def test_source_image_conversion_uses_wps_pdf_then_bundled_rasterizer(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """非 Windows WPS 不提供 EnhMetaFileBits，页面布局仍必须由源码/WPS 生成。"""

    monkeypatch.delenv("PARTYOPS_FORMATTER_TEST_LOCAL_ENGINE", raising=False)
    source = tmp_path / "source.docx"
    Document().save(source)
    source_pdf = tmp_path / "source-wps.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page(width=595, height=842)
        page.insert_text((72, 96), "PartyOps WPS page")
        pdf.save(source_pdf)
    calls: list[tuple[str, Path, Path, dict[str, object]]] = []

    def fake_host(
        feature_id: str,
        source_path: Path,
        workspace: Path,
        options: dict[str, object],
        **_kwargs: object,
    ) -> tuple[official_format_host.SourceHostOutput, ...]:
        calls.append((feature_id, source_path, workspace, options))
        return (
            official_format_host.SourceHostOutput(
                path=source_pdf,
                host_display_name="WPS Office",
                message="PDF 转换完成",
            ),
        )

    monkeypatch.setattr(official_format_host, "run_source_host", fake_host)
    result = official_format_features.execute_feature(
        "convert",
        source,
        tmp_path / "job",
        {"target_format": "png", "image_mode": "pages", "dpi": 144},
    )
    assert len(calls) == 1
    assert calls[0][:2] == ("convert", source)
    assert calls[0][2] == tmp_path / "job" / "source-pdf"
    assert calls[0][3]["target_format"] == "pdf"
    assert len(result.outputs) == 1
    assert result.outputs[0].path.suffix == ".png"
    assert result.outputs[0].path.is_file()
