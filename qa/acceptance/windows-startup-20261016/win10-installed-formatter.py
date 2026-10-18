"""原 Win10 x64 安装版排版薄适配；仅显式 --execute 才连接 Guest。"""
from __future__ import annotations

import argparse
import http.cookiejar
import http.client
import io
import importlib.util
import json
import ntpath
import re
import sys
import struct
import urllib.parse
import urllib.error
import urllib.request
import uuid
from pathlib import Path
from types import SimpleNamespace

import retention_contract as contract

ROOT = Path(__file__).resolve().parents[3]
OUTPUT = Path(__file__).resolve().parent
LAB_ROOT = Path("D:/PartyOps-VM-Lab")
REMOTE = "lifecycle-82dbfdf1894f"
WPS_VERSION = "12.1.0.28488"
# 来自本次官方28488包实际安装文件；不同用户安装必须仍是相同程序字节。
WPS_EXECUTABLE_SHA256 = "4ea193b0801a88752b69c4281cfe7513a3ff945afd57aca59fe13f15a87d9798"
FZX_FONT = Path("C:/Users/Administrator/AppData/Local/Microsoft/Windows/Fonts/方正小标宋简.TTF")
FZX_FONT_SHA256 = "f6154b97ee71bad19efa7f78edc152cf515c1f943220fc6e98360005d53d1c62"
FONT_WITNESSES = (
    ("宋体", "字体验收甲", ("宋体", "SimSun")),
    ("黑体", "字体验收乙", ("黑体", "SimHei")),
    ("仿宋_GB2312", "字体验收丙", ("仿宋_GB2312", "FangSong_GB2312")),
    ("楷体_GB2312", "字体验收丁", ("楷体_GB2312", "KaiTi_GB2312")),
    ("方正小标宋简体", "字体验收戊", ("方正小标宋简体", "FZXiaoBiaoSong-B05S")),
)
sys.path.insert(0, str(ROOT / "qa/vm-lab"))


def load(path, name):
    spec = importlib.util.spec_from_file_location(name, path)
    result = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(result)
    return result


def require(value, code):
    if not value:
        raise RuntimeError(code)


def safe_job(value):
    """通用 runner 只收到作业断言必要字段；消息、路径、令牌均不进入文件。"""
    require(isinstance(value, dict) and re.fullmatch(r"[a-f0-9]{32}", str(value.get("id", ""))), "JOB_ID_INVALID")
    require(value.get("state") in {"queued", "running", "completed", "failed", "cancelled", "completed_with_errors"}, "JOB_STATE_INVALID")
    outputs = []
    for item in value.get("outputs", []):
        require(re.fullmatch(r"[a-f0-9]{32}", str(item.get("id", ""))), "OUTPUT_ID_INVALID")
        require(re.fullmatch(r"[a-f0-9]{32}", str(item.get("document_id", ""))), "OUTPUT_DOCUMENT_ID_INVALID")
        filename = ntpath.basename(str(item.get("filename", item.get("name", ""))))
        require(filename and len(filename) < 200 and not any(c in filename for c in "\r\n"), "OUTPUT_NAME_INVALID")
        outputs.append({"id": item["id"], "document_id": item.get("document_id"), "filename": filename})
    return {"id": value["id"], "state": value["state"], "outputs": outputs}


def route_template(path):
    """只保留已知路由词；参数、查询串和未知片段均不落盘。"""
    known = {"api", "v1", "health", "auth", "login", "official-format", "local-ticket", "sessions",
             "capabilities", "documents", "jobs", "outputs", "cancel"}
    parts = urllib.parse.urlsplit(path).path.split("/")
    return "/".join(part if part in known or not part else "{id}" for part in parts)


def http_failure(error, path, method, status=None):
    """读取错误响应仅提取协议错误码；正文、头、URL和异常消息均不保存。"""
    row = {"route": route_template(path), "method": method if method in {"GET", "POST", "DELETE", "PUT", "PATCH"} else "UNKNOWN",
           "exception_type": type(error).__name__, "http_status": status if type(status) is int and 100 <= status <= 599 else None}
    if isinstance(error, http.client.IncompleteRead):
        # 仅计数；partial含响应正文，不得落盘或拼入异常消息。
        row["partial_bytes"] = len(error.partial)
        row["expected_remaining_bytes"] = error.expected if type(error.expected) is int and error.expected >= 0 else None
    if isinstance(error, urllib.error.HTTPError):
        row["http_status"] = error.code if type(error.code) is int and 100 <= error.code <= 599 else None
        try:
            value = json.loads(error.read(65536))
            code = value.get("error", {}).get("code") if isinstance(value, dict) and isinstance(value.get("error"), dict) else None
            if isinstance(code, str) and re.fullmatch(r"[A-Z][A-Z0-9_]{0,99}", code):
                row["error_code"] = code
        except Exception:
            pass
    return row


def session_transport_metadata(request, response, payload):
    """仅记录新会话创建和DELETE的安全计数，不保存请求/响应正文或头值。"""
    method = request.get_method()
    path = urllib.parse.urlsplit(request.full_url).path
    if not ((method == "POST" and path == "/v1/sessions")
            or (method == "DELETE" and re.fullmatch(r"/v1/sessions/[a-f0-9]{32}", path))):
        return None
    decimal = lambda value: int(value) if isinstance(value, str) and re.fullmatch(r"[0-9]{1,12}", value) else None
    request_length = decimal(request.get_header("Content-length"))
    if request_length is None and request.data is None and method == "POST":
        request_length = 0  # http.client为无body POST自动发送Content-Length: 0。
    return {"route": route_template(path), "method": method, "status": response.status,
            "protocol": response.version, "request_body_bytes": 0 if request.data is None else len(request.data),
            "request_content_type_present": any(key.lower() == "content-type" for key, _ in request.header_items()),
            "request_content_length": request_length,
            "response_content_length": decimal(response.headers.get("Content-Length")), "actual_response_bytes": len(payload)}


def save_failure_evidence(output, record, jobs, probes):
    """在删除Guest会话前写出已经取得的证据，不重新读取或执行作业。"""
    record["safe_jobs"] = list(jobs.values())
    record["probe_sidecars"] = []
    for directory in probes:
        if re.fullmatch(r"win10-formatter-[a-f0-9]{12}", directory):
            sidecar = OUTPUT / (directory + ".probe.json")
            if sidecar.is_file():
                # 内容由staged_readonly_probe白名单产生；保留一份在本次run目录，清理前可定位。
                (output / sidecar.name).write_bytes(sidecar.read_bytes())
                record["probe_sidecars"].append(sidecar.name)
    snapshot = {key: record[key] for key in ("stage", "error", "http_failures", "session_transport", "last_request", "safe_jobs", "probe_sidecars") if key in record}
    (output / "failure-evidence.json").write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")


def listener_script(port):
    require(isinstance(port, int) and 1 <= port <= 65535, "LISTENER_PORT_INVALID")
    return f"$port={port}\n" + r"""
$ErrorActionPreference='Stop'
$rows=@(Get-NetTCPConnection -State Listen -LocalPort $port -ErrorAction SilentlyContinue)
$owners=@($rows|Select-Object -ExpandProperty OwningProcess -Unique)
if($owners.Count -ne 1 -or @($rows|Where-Object {$_.LocalAddress -ne '127.0.0.1'}).Count) {throw 'LOOPBACK_LISTENER_REQUIRED'}
$p=Get-CimInstance Win32_Process -Filter ('ProcessId='+$owners[0])
$owner=Invoke-CimMethod -InputObject $p -MethodName GetOwner
if($owner.ReturnValue -ne 0 -or $owner.Domain -ne $env:COMPUTERNAME) {throw 'LOCAL_PROCESS_OWNER_REQUIRED'}
@{pid=[int]$p.ProcessId;path=$p.ExecutablePath;created=$p.CreationDate.ToUniversalTime().ToString('o');owner=$owner.User;session_id=[int]$p.SessionId}|ConvertTo-Json -Compress
"""


def validate_listener(row, actual, desktop):
    process = actual["processes"][0]
    require(row.get("pid") == process["pid"] and row.get("created") == process["created"]
            and row.get("path", "").casefold() == (contract.INSTALL + r"\PartyOps.exe").casefold()
            and row.get("owner") == "partyopsuser" and row.get("session_id") == desktop["desktop"][0]["session_id"],
            "FORMATTER_LISTENER_OWNER_CHANGED")


def normalize_wps_version(value):
    parts = re.fullmatch(r"\s*(\d+)\s*[.,]\s*(\d+)\s*[.,]\s*(\d+)\s*[.,]\s*(\d+)\s*", str(value))
    require(parts is not None, "WPS_NUMERIC_VERSION_INVALID")
    return ".".join(str(int(part)) for part in parts.groups())


def validate_wps(value):
    require(normalize_wps_version(value.get("version")) == WPS_VERSION and value.get("pe_machine") == 0x8664
            and value.get("com_registered") is True
            and value.get("executable_sha256") == WPS_EXECUTABLE_SHA256
            and value.get("signature_status") == "Valid" and value.get("signer_is_kingsoft") is True,
            "WPS_X64_DETECTION_INVALID")
    exe = ntpath.normcase(ntpath.normpath(value.get("executable_path", "")))
    user_root = ntpath.normcase(r"C:\Users\partyopsuser\AppData\Local\Kingsoft") + "\\"
    shared_root = ntpath.normcase(r"C:\Program Files\Kingsoft") + "\\"
    require(value.get("username") == "partyopsuser"
            and value.get("registration_scope") in {"target-user", "machine"}
            and value.get("identity_sources") and set(value["identity_sources"]) <= {"target-user-vendor", "machine-vendor", "target-user-app-path", "machine-app-path"}
            and ((value.get("binary_scope") == "target-user-profile" and exe.startswith(user_root))
                 or (value.get("binary_scope") == "shared-program-files" and exe.startswith(shared_root)))
            and ntpath.basename(exe) == "wps.exe", "WPS_ORIGINAL_USER_SCOPE_REQUIRED")
    fonts = value.get("fonts", [])
    require(len(fonts) == 5 and {row.get("font") for row in fonts} ==
            {"宋体", "黑体", "仿宋_GB2312", "楷体_GB2312", "方正小标宋简体"}, "FONT_READ_PROOF_INCOMPLETE")
    require(all(row.get("file_load_readable") is True for row in fonts), "REQUIRED_FONT_FILES_NOT_READABLE")


def readonly_probe_command(directory):
    require(re.fullmatch(r"win10-formatter-[a-f0-9]{12}", directory), "QA_SCRIPT_DIRECTORY_INVALID")
    remote = "C:/ProgramData/PartyOps-VM-Lab/" + directory
    return "& ([scriptblock]::Create([IO.File]::ReadAllText('" + remote + "/wps-readonly.ps1',[Text.Encoding]::UTF8)))"


def staged_readonly_probe(client, execute, script, directory, binding=None):
    """长只读脚本走既有 SFTP，避免 SSH/cmd EncodedCommand 长度限制。"""
    command = readonly_probe_command(directory)
    remote = "C:/ProgramData/PartyOps-VM-Lab/" + directory
    # 仅回传本脚本显式声明的断言码；系统异常、路径及远端 stderr 均不出 Guest。
    codes = sorted(set(re.findall(r"throw '([A-Z][A-Z0-9_]+)'", script)))
    failure = "WPS_X64_READONLY_PROBE_FAILED"
    allowed = ",".join("'" + code + "'" for code in codes)
    wrapped = ("try { & {\n" + script + "\n} } catch {\n"
               "$code=[string]$_.Exception.Message\n"
               "if($code -cnotin @(" + allowed + ")){$code='" + failure + "'}\n"
               "@{qa_probe_error=$code;exception_type=$_.Exception.GetType().FullName;"
               "category=$_.CategoryInfo.Category.ToString();script_line=[int]$_.InvocationInfo.ScriptLineNumber}"
               "|ConvertTo-Json -Compress\n}")
    def fail(code, diagnostic=None):
        # 失败证据只允许类型、枚举、行号和合成作业摘要；不保存正文、路径或 stderr。
        diagnostic = diagnostic if isinstance(diagnostic, dict) else {}
        row = {"code": code, "exception_type": "Unknown", "category": "NotSpecified", "script_line": 0}
        kind = diagnostic.get("exception_type")
        if isinstance(kind, str) and re.fullmatch(r"(?:System\.)?[A-Za-z][A-Za-z0-9_.]{0,150}", kind):
            row["exception_type"] = kind
        categories = {"NotSpecified", "InvalidData", "InvalidArgument", "InvalidOperation", "ObjectNotFound",
                      "PermissionDenied", "ReadError", "WriteError", "ParserError", "OperationStopped",
                      "ResourceUnavailable", "OpenError", "MetadataError", "LimitsExceeded", "SecurityError",
                      "ConnectionError", "ProtocolError", "OperationTimeout", "SyntaxError", "NotImplemented",
                      "InvalidResult", "DeviceError", "FromStdErr", "ResourceBusy", "ResourceExists", "NotEnabled",
                      "AuthenticationError", "CloseError", "DeadlockDetected", "QuotaExceeded"}
        if isinstance(diagnostic.get("category"), str) and diagnostic["category"] in categories:
            row["category"] = diagnostic["category"]
        line = diagnostic.get("script_line")
        if type(line) is int and 0 <= line <= len(wrapped.splitlines()):
            row["script_line"] = line
        row["binding"] = {}
        for key, pattern in (("job", r"[a-f0-9]{32}"), ("feature", r"format|replace|redheader|rename|convert|pdf-to-word"),
                             ("input_sha256", r"[a-f0-9]{64}"), ("host_sha256", r"[a-f0-9]{64}")):
            item = (binding or {}).get(key)
            if isinstance(item, str) and re.fullmatch(pattern, item):
                row["binding"][key] = item
        contract.exclusive(OUTPUT / (directory + ".probe.json"), row)
        raise RuntimeError(code) from None
    try:
        with client.open_sftp() as fs:
            fs.mkdir(remote)
            with fs.file(remote + "/wps-readonly.ps1", "wx") as stream:
                stream.write(wrapped.encode("utf-8"))
        result = json.loads(execute(client, command, timeout=240))
    except Exception as exc:
        fail(failure, {"exception_type": type(exc).__name__, "category": "InvalidData" if isinstance(exc, json.JSONDecodeError) else "ConnectionError"})
    if isinstance(result, dict) and "qa_probe_error" in result:
        code = result["qa_probe_error"]
        fail(code if isinstance(code, str) and code in codes else failure, result)
    if not isinstance(result, dict):
        fail(failure, {"category": "InvalidData"})
    return result


def wps_script(sid):
    require(re.fullmatch(r"S-1-5-21-(?:\d+-){3}\d+", str(sid)), "STANDARD_USER_SID_INVALID")
    return "$sid='" + sid + "'\n" + r"""
$ErrorActionPreference='Stop'
$account=Get-LocalUser -Name 'partyopsuser'
if($account.SID.Value -ne $sid -or -not $account.Enabled -or @(Get-LocalGroupMember -SID 'S-1-5-32-544'|Where-Object {$_.SID.Value -eq $sid}).Count){throw 'WPS_TARGET_STANDARD_USER_REQUIRED'}
# 厂商实际不写InstallLocation；只接受目标用户/机器登记，不搜索其它用户或递归猜目录。
$registrations=@(
  @{root=('Registry::HKEY_USERS\'+$sid+'\Software');scope='target-user'},
  @{root='HKLM:\SOFTWARE';scope='machine'}
)
$candidates=@(foreach($registration in $registrations){
  $vendor=Join-Path $registration.root 'Kingsoft\Office\6.0\Common'
  if(Test-Path -LiteralPath $vendor){
    $installRoot=[string](Get-Item -LiteralPath $vendor).GetValue('InstallRoot')
    if($installRoot -and [IO.Path]::IsPathRooted($installRoot) -and $installRoot -notmatch '%'){
      $path=Join-Path $installRoot.Trim('"') 'office6\wps.exe'
      if(Test-Path -LiteralPath $path -PathType Leaf){@{path=[IO.Path]::GetFullPath($path);source=($registration.scope+'-vendor')}}
    }
  }
  $appPath=Join-Path $registration.root 'Microsoft\Windows\CurrentVersion\App Paths\wps.exe'
  if(Test-Path -LiteralPath $appPath){
    $path=([string](Get-Item -LiteralPath $appPath).GetValue('')).Trim('"')
    if([IO.Path]::IsPathRooted($path) -and $path -notmatch '%' -and (Test-Path -LiteralPath $path -PathType Leaf)){
      @{path=[IO.Path]::GetFullPath($path);source=($registration.scope+'-app-path')}
    }
  }
})
$paths=@($candidates|ForEach-Object {$_.path}|Sort-Object -Unique)
if($paths.Count -ne 1){throw 'WPS_VENDOR_EXECUTABLE_MISSING_OR_AMBIGUOUS'}
$exe=Get-Item -LiteralPath $paths[0]
if($exe.FullName.StartsWith('C:\Users\partyopsuser\AppData\Local\Kingsoft\',[StringComparison]::OrdinalIgnoreCase)){$binaryScope='target-user-profile'}
elseif($exe.FullName.StartsWith('C:\Program Files\Kingsoft\',[StringComparison]::OrdinalIgnoreCase)){$binaryScope='shared-program-files'}
else{throw 'WPS_OTHER_USER_OR_UNOBSERVED_SCOPE_REJECTED'}
$signature=Get-AuthenticodeSignature -LiteralPath $exe.FullName
if($signature.Status.ToString() -ne 'Valid' -or $signature.SignerCertificate.Subject -notmatch 'Kingsoft'){throw 'WPS_KINGSOFT_SIGNATURE_REQUIRED'}
$exeHash=(Get-FileHash -LiteralPath $exe.FullName -Algorithm SHA256).Hash.ToLowerInvariant()
if($exeHash -ne '4ea193b0801a88752b69c4281cfe7513a3ff945afd57aca59fe13f15a87d9798'){throw 'WPS_OBSERVED_EXECUTABLE_HASH_CHANGED'}
$r=New-Object IO.BinaryReader([IO.File]::OpenRead($exe.FullName))
try {
  if($r.ReadUInt16() -ne 23117) {throw 'WPS_PE_INVALID'}
  $r.BaseStream.Position=60;$offset=$r.ReadUInt32()
  if($offset -lt 64 -or $offset -gt $r.BaseStream.Length-26) {throw 'WPS_PE_INVALID'}
  $r.BaseStream.Position=$offset
  if($r.ReadUInt32() -ne 17744 -or $r.ReadUInt16() -ne 34404) {throw 'WPS_X64_EXECUTABLE_REQUIRED'}
  $r.BaseStream.Position=$offset+24;if($r.ReadUInt16() -ne 523) {throw 'WPS_PE32PLUS_REQUIRED'}
}finally {$r.Close()}
$version=@($exe.VersionInfo.FileMajorPart,$exe.VersionInfo.FileMinorPart,$exe.VersionInfo.FileBuildPart,$exe.VersionInfo.FilePrivatePart)-join '.'
if($version -ne '12.1.0.28488'){throw 'WPS_X64_VERSION_MISMATCH'}
$classes=@(('Registry::HKEY_USERS\'+$sid+'\Software\Classes'),'HKLM:\SOFTWARE\Classes')
$registrationScope='machine';$clsid=$null
# 模拟目标用户HKCR优先级；目标用户错误覆盖不能由机器登记冒充成功。
for($i=0;$i -lt $classes.Count;$i++){
  $key=Join-Path $classes[$i] 'KWPS.Application\CLSID'
  if(Test-Path -LiteralPath $key){$clsid=[string](Get-Item -LiteralPath $key).GetValue('');if($i -eq 0){$registrationScope='target-user'};break}
}
if($clsid -notmatch '^\{[a-fA-F0-9-]{36}\}$'){throw 'WPS_USER_COM_CLASS_MISSING'}
$serverValue=$null
for($i=0;$i -lt $classes.Count;$i++){
  $key=Join-Path $classes[$i] ('CLSID\'+$clsid+'\LocalServer32')
  if(Test-Path -LiteralPath $key){$serverValue=[string](Get-Item -LiteralPath $key).GetValue('');if($i -eq 0){$registrationScope='target-user'};break}
}
$serverPath=$null
if($serverValue -match '^"([^"]+\.exe)"(?:\s.*)?$'){$serverPath=$Matches[1]}
elseif($serverValue -match '^(.+?\.exe)(?:\s.*)?$'){$serverPath=$Matches[1]}
if(-not $serverPath -or [IO.Path]::GetFullPath($serverPath) -ine $exe.FullName){throw 'WPS_TARGET_USER_COM_EXECUTABLE_MISMATCH'}
# 实际加载登记的字体文件并读取字体族；不安装、不改登记，仅输出目标字体读取状态。
Add-Type -AssemblyName System.Drawing
$fontFiles=@(foreach($root in @('HKLM:\SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts',('Registry::HKEY_USERS\'+$sid+'\Software\Microsoft\Windows NT\CurrentVersion\Fonts'))) {
  if(Test-Path -LiteralPath $root) {$key=Get-Item -LiteralPath $root
    foreach($name in $key.GetValueNames()) {$path=[string]$key.GetValue($name)
      if(-not [IO.Path]::IsPathRooted($path)) {$path=Join-Path ($env:WINDIR+'\Fonts') $path}
      if(Test-Path -LiteralPath $path -PathType Leaf) {$path}}}
})|Sort-Object -Unique
$families=@(foreach($path in $fontFiles) {$collection=New-Object Drawing.Text.PrivateFontCollection
  try {$collection.AddFontFile($path);foreach($family in $collection.Families) {$family.Name}}catch {}finally {$collection.Dispose()}
})|Sort-Object -Unique
$fonts=@(foreach($pair in @(@('宋体','SimSun'),@('黑体','SimHei'),@('仿宋_GB2312','FangSong_GB2312'),@('楷体_GB2312','KaiTi_GB2312'),@('方正小标宋简体','FZXiaoBiaoSong-B05S'))) {
  @{font=$pair[0];file_load_readable=(@($families|Where-Object {$_ -in $pair}).Count -gt 0)}
})
@{version=$version;pe_machine=34404;executable_sha256=$exeHash;signature_status=$signature.Status.ToString();signer_is_kingsoft=$true;username='partyopsuser';binary_scope=$binaryScope;registration_scope=$registrationScope;identity_sources=@($candidates|ForEach-Object {$_.source}|Sort-Object -Unique);executable_path=$exe.FullName;com_registered=$true;fonts=$fonts}|ConvertTo-Json -Depth 5 -Compress
"""


def font_witness(path):
    """只生成合成QA文档；字体名称声明本身不计渲染通过。"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    document = Document()
    for font, marker, _ in FONT_WITNESSES:
        run = document.add_paragraph().add_run(marker)
        run.font.name, run.font.size = font, Pt(20)
        names = run._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            names.set(qn("w:" + key), font)
    document.save(path)


def normalized_pdf_font(value):
    return re.sub(r"[\s_-]", "", re.sub(r"^[A-Z]{6}\+", "", value)).casefold()


def font_trace_proof(traces, witnesses=FONT_WITNESSES):
    """只接受PDF可见绘制trace里的实际字体和对应合成字形；拒绝字体资源表/文档声明。"""
    result = []
    for font, marker, aliases in witnesses:
        names = {normalized_pdf_font(name) for name in aliases}
        selected = [s for s in traces if normalized_pdf_font(s.get("font", "")) in names
                    and s.get("type") in {0, 1, 2} and s.get("opacity", 0) > 0
                    and s.get("size", 0) > 0 and len(s.get("bbox", [])) == 4
                    and s["bbox"][2] > s["bbox"][0] and s["bbox"][3] > s["bbox"][1]]
        text = "".join("".join(chr(char[0]) for char in span.get("chars", [])
                              if len(char) >= 2 and char[0] > 0 and char[1] > 0) for span in selected)
        require(marker in text, "PDF_ACTUAL_FONT_GLYPHS_NOT_PROVEN")
        result.append({"font": font, "actual_pdf_fonts": sorted({s["font"] for s in selected}),
                       "synthetic_marker_drawn": True, "visible_trace_count": len(selected)})
    return result


def pinned_font_embedding_type():
    """读取固定QA字体的OS/2字段；不改变字体字节或嵌入许可。"""
    require(contract.digest(FZX_FONT) == FZX_FONT_SHA256, "FZX_REFERENCE_FONT_CHANGED")
    data = FZX_FONT.read_bytes()
    count = struct.unpack_from(">H", data, 4)[0]
    for index in range(count):
        tag, _, offset, length = struct.unpack_from(">4sIII", data, 12 + index * 16)
        if tag == b"OS/2" and length >= 10:
            value = struct.unpack_from(">H", data, offset + 8)[0]
            require(value == 2, "FZX_REFERENCE_EMBEDDING_TYPE_CHANGED")
            return value
    raise RuntimeError("FZX_REFERENCE_OS2_REQUIRED")


def raster_glyph_metrics(mask, reference):
    """独立字形渲染与PDF透明蒙版逐像素比较；仅统一ink边界和采样尺寸。"""
    from PIL import Image, ImageChops, ImageStat
    require(mask.mode == reference.mode == "L" and mask.getbbox() and reference.getbbox(), "FONT_RASTER_INK_REQUIRED")
    actual = mask.crop(mask.getbbox())
    expected = reference.crop(reference.getbbox()).resize(actual.size, Image.Resampling.LANCZOS)
    binary = lambda image: image.point(lambda value: 255 if value >= 128 else 0).convert("1")
    def iou(first, second):
        left, right = binary(first), binary(second)
        union = ImageChops.logical_or(left, right).histogram()[255]
        return ImageChops.logical_and(left, right).histogram()[255] / union if union else 0
    return {"iou": iou(actual, expected), "mae": ImageStat.Stat(ImageChops.difference(actual, expected)).mean[0] / 255,
            "per_glyph_iou": [iou(actual.crop((round(actual.width * i / 5), 0, round(actual.width * (i + 1) / 5), actual.height)),
                                  expected.crop((round(actual.width * i / 5), 0, round(actual.width * (i + 1) / 5), actual.height)))
                              for i in range(5)], "ink_pixels": list(actual.size)}


def require_raster_glyph_match(metrics):
    require(metrics["iou"] >= 0.94 and metrics["mae"] <= 0.025 and len(metrics["per_glyph_iou"]) == 5
            and min(metrics["per_glyph_iou"]) >= 0.90, "FZX_RASTER_GLYPHS_NOT_PROVEN")


def font_render_proof(pdf_path):
    """可嵌入字体仍核texttrace；固定受限小标宋另核独立TTF渲染的可见位图字形来源。"""
    import fitz
    from PIL import Image, ImageDraw, ImageFont, ImageOps
    with fitz.open(pdf_path) as pdf:
        traces = [span for page in pdf for span in page.get_texttrace()]
        try:
            return font_trace_proof(traces)
        except RuntimeError:
            first = font_trace_proof(traces, FONT_WITNESSES[:4])
        embedding = pinned_font_embedding_type()
        require(len(pdf) == 1, "FZX_RASTER_SINGLE_PAGE_REQUIRED")
        page = pdf[0]
        images = page.get_images()
        require(len(images) == 1 and images[0][1] > 0, "FZX_VISIBLE_SMASK_REQUIRED")
        xref, smask, width, height = images[0][:4]
        rectangles = page.get_image_rects(xref)
        log = page.get_bboxlog()
        require(len(rectangles) == 1 and log and log[-1][0] == "fill-image"
                and all(abs(a - b) < 0.001 for a, b in zip(rectangles[0], log[-1][1]))
                and not rectangles[0].is_empty and page.rect.contains(rectangles[0])
                and rectangles[0].y0 >= max(span["bbox"][3] for span in traces), "FZX_RASTER_VISIBLE_PLACEMENT_REQUIRED")
        pixmap = fitz.Pixmap(pdf, smask)
        require(pixmap.n == 1 and (pixmap.width, pixmap.height) == (width, height), "FZX_SMASK_DIMENSIONS_INVALID")
        mask = Image.frombytes("L", (width, height), pixmap.samples)
        reference = Image.new("L", (2500, 1000))
        ImageDraw.Draw(reference).text((50, 50), FONT_WITNESSES[-1][1], fill=255,
                                       font=ImageFont.truetype(str(FZX_FONT), 400))
        metrics = raster_glyph_metrics(mask, reference)
        require_raster_glyph_match(metrics)
        visible = page.get_pixmap(matrix=fitz.Matrix(20, 20), clip=rectangles[0], alpha=False)
        visible_image = ImageOps.invert(Image.frombytes("RGB", (visible.width, visible.height), visible.samples).convert("L"))
        visible_metrics = raster_glyph_metrics(visible_image, reference)
        require_raster_glyph_match(visible_metrics)
        return first + [{"font": FONT_WITNESSES[-1][0], "synthetic_marker_drawn": True,
                         "actual_font_source": "pinned-ttf-independent-glyph-raster-match",
                         "reference_sha256": FZX_FONT_SHA256, "reference_os2_fs_type": embedding,
                         "renderer": "Pillow-FreeType", "reference_pixel_size": 400, "metrics": metrics,
                         "visible_page_metrics": visible_metrics, "pdf_behavior": "restricted-embedding-font-rasterized",
                         "pdf_image_xref": xref, "pdf_smask_xref": smask, "visible_bbox": list(rectangles[0])}]


def resume_font_evidence(folder, binding, host_sha, oracle):
    """续验仅复用同候选已完成的九功能与金样；重核每个下载摘要、原生补证和本地金样。"""
    folder = Path(folder).resolve()
    require(folder.parent == OUTPUT and re.fullmatch(r"win10-formatter-[a-f0-9]{12}", folder.name), "FONT_RESUME_DIRECTORY_INVALID")
    summary_path = folder / "summary.json"
    previous = contract.read(summary_path)
    require(previous.get("candidate") == binding and previous.get("status") == "failed"
            and previous.get("error") == "PDF_ACTUAL_FONT_GLYPHS_NOT_PROVEN" and previous.get("failure_stage") == "font-render",
            "FONT_RESUME_SOURCE_BINDING_INVALID")
    checks = {row["id"]: row for row in previous["checks"]}
    feature_path = folder / "installed-features-evidence.json"
    features = contract.read(feature_path)
    manifest_path = folder / "features-manifest.json"
    require(features.get("manifest_sha256") == contract.digest(manifest_path), "FONT_RESUME_FEATURE_MANIFEST_CHANGED")
    manifest = {row["id"]: row for row in contract.read(manifest_path)["cases"]}
    verify_native_cases(features, checks["feature-native-source-host"]["proofs"], host_sha)
    require(features.get("case_count") == 9 and checks["six-feature-http-output-contracts"].get("status") == "passed",
            "FONT_RESUME_NINE_FEATURES_REQUIRED")
    for case in features["cases"]:
        require(case.get("status") == "passed", "FONT_RESUME_FEATURE_NOT_PASSED")
        original = manifest.get(case["id"], {})
        require(original.get("feature_id") == case["feature_id"]
                and checks["feature-native-source-host"]["proofs"][case["job"]["id"]].get("input_sha256") == original.get("sha256"),
                "FONT_RESUME_FEATURE_INPUT_BINDING_CHANGED")
        for item in case["outputs"]:
            path = folder / item["filename"]
            require(path.resolve().parent == folder and contract.digest(path) == item["sha256"], "FONT_RESUME_FEATURE_DOWNLOAD_CHANGED")
    golden = checks["format-and-pdf-golden"]
    require(golden.get("status") == "passed" and len(golden.get("host_proofs", [])) == 5
            and len(golden.get("visual_comparisons", [])) == 2 and all(row["passed"] for row in golden["visual_comparisons"]),
            "FONT_RESUME_GOLDEN_REQUIRED")
    proof_key = lambda row: tuple(row.get(key) for key in ("input_sha256", "request_sha256", "response_sha256"))
    wanted = {proof_key(row) for row in golden["host_proofs"]}
    require(len(wanted) == 5, "FONT_RESUME_GOLDEN_PROOF_SET_INVALID")
    required_names = {"formatted-1.docx", "formatted-2.docx", "render-0.pdf", "render-1.pdf", "render-2.pdf"}
    seen, names, sources = set(), set(), []
    for supplement in sorted(folder.glob("native-output-supplement-*.json")):
        if not re.fullmatch(r"native-output-supplement-[a-f0-9]{32}-[a-f0-9]{12}\.json", supplement.name):
            continue
        row = contract.read(supplement)
        require(re.fullmatch(r"[a-f0-9]{32}", str(row.get("job", "")))
                and contract.digest(folder / ("job-" + row["job"] + ".json")) == row.get("original_job_sha256"),
                "FONT_RESUME_SUPPLEMENT_JOB_CHANGED")
        relevant = False
        for item in row.get("outputs", []):
            proof = item["native_proof"]
            key = proof_key(proof)
            if key not in wanted:
                continue
            path = folder / item["file"]
            native = proof.get("outputs", [])
            require(item["file"] in required_names and path.resolve().parent == folder and len(native) == 1
                    and contract.digest(path) == item["sha256"] == native[0].get("sha256")
                    and path.stat().st_size == item["bytes"] == native[0].get("bytes")
                    and path.suffix.lower() == native[0].get("extension"), "FONT_RESUME_GOLDEN_DOWNLOAD_CHANGED")
            seen.add(key); names.add(item["file"]); relevant = True
        if relevant:
            sources.append({"file": supplement.name, "sha256": contract.digest(supplement)})
    require(seen == wanted and names == required_names, "FONT_RESUME_NATIVE_GOLDEN_BINDING_REQUIRED")
    source = ROOT / "backend/tests/fixtures/document-formatter-source/expected-source-formatted.docx"
    signature = oracle._semantic_equivalence_signature(oracle._document_signature(source))
    for name in ("formatted-1.docx", "formatted-2.docx"):
        document = oracle._document_signature(folder / name)
        oracle._require_contract(document)
        require(oracle._semantic_equivalence_signature(document) == signature, "FONT_RESUME_GOLDEN_SEMANTIC_CHANGED")
    for name in ("render-1.pdf", "render-2.pdf"):
        require(oracle._visual_page_comparison(folder / "render-0.pdf", folder / name, allow_rasterization_noise=False)["passed"],
                "FONT_RESUME_GOLDEN_PIXELS_CHANGED")
    return [checks[key] for key in ("feature-native-source-host", "six-feature-http-output-contracts", "format-and-pdf-golden")], {
        "run": folder.name, "summary_sha256": contract.digest(summary_path), "features_sha256": contract.digest(feature_path),
        "features_manifest_sha256": contract.digest(manifest_path),
        "supplements": sources, "validation": "nine-feature-downloads-native-proofs-and-golden-semantics-pixels-rechecked"}


def native_host_script(temp, job, document, feature, input_sha, host_sha):
    """沿用正式Win10 x86 host_proof的当前随机job查找；扩展到六入口并核下载输出摘要。"""
    require(re.fullmatch(r"[A-Za-z]:\\Users\\[^\\'\r\n]+\\AppData\\Local\\Temp", temp), "WPS_USER_TEMP_PATH_INVALID")
    require(all(re.fullmatch(r"[a-f0-9]{32}", v) for v in (job, document))
            and all(re.fullmatch(r"[a-f0-9]{64}", v) for v in (input_sha, host_sha))
            and feature in {"format", "replace", "redheader", "rename", "convert", "pdf-to-word"}, "NATIVE_PROOF_ARGUMENT_INVALID")
    value = {"temp": temp, "job": job, "document": document, "feature": feature, "input_sha": input_sha,
             "host_sha": host_sha, "host": contract.INSTALL + r"\formatter-host\PartyOps.DocumentFormatter.Host.exe"}
    return "$expected='" + json.dumps(value).replace("'", "''") + "'|ConvertFrom-Json\n" + r"""
$ErrorActionPreference='Stop'
if((Get-FileHash -LiteralPath $expected.host -Algorithm SHA256).Hash.ToLowerInvariant() -ne $expected.host_sha){throw 'CANDIDATE_NATIVE_HOST_CHANGED'}
$root=[IO.Path]::GetFullPath($expected.temp).TrimEnd('\')+'\'
$matches=@(foreach($folder in @(Get-ChildItem -LiteralPath $expected.temp -Directory)){
  if($folder.Attributes -band [IO.FileAttributes]::ReparsePoint){continue}
  if($folder.Name -like 'partyops-official-format-*'){$parts=@('jobs',$expected.job,$expected.document)}
  elseif($folder.Name -like 'pf-*'){$parts=@('j',$expected.job,'0')}else{continue}
  $parent=$folder.FullName;foreach($part in $parts){$parent=Join-Path $parent $part}
  foreach($suffix in @('.source-host','source-pdf\.source-host')){
    $dir=Join-Path $parent $suffix
    if(Test-Path -LiteralPath (Join-Path $dir 'response.json') -PathType Leaf){$dir}
  }
})
if($matches.Count -ne 1){throw 'CURRENT_NATIVE_HOST_PROOF_AMBIGUOUS'}
$requestPath=Join-Path $matches[0] 'request.json';$responsePath=Join-Path $matches[0] 'response.json'
# Python和原生宿主写UTF-8；PS5默认ANSI不能读取无BOM中文路径。
$request=Get-Content -LiteralPath $requestPath -Raw -Encoding UTF8|ConvertFrom-Json
$response=Get-Content -LiteralPath $responsePath -Raw -Encoding UTF8|ConvertFrom-Json
if($request.feature_id -ne $expected.feature -or $request.host_preference -ne 'wps' -or @($request.source_paths).Count -ne 1){throw 'CURRENT_NATIVE_REQUEST_MISMATCH'}
$source=[IO.Path]::GetFullPath([string]$request.source_paths[0])
if(-not $source.StartsWith($root,[StringComparison]::OrdinalIgnoreCase) -or (Get-FileHash -LiteralPath $source -Algorithm SHA256).Hash.ToLowerInvariant() -ne $expected.input_sha){throw 'CURRENT_NATIVE_INPUT_MISMATCH'}
$jobs=@($response.jobs)
if($response.fatal -or $jobs.Count -ne 1 -or $jobs[0].success -ne $true){throw 'CURRENT_NATIVE_HOST_FAILED'}
$provider=[string]$jobs[0].host_display_name
if($expected.feature -ne 'pdf-to-word' -and $provider -notmatch 'WPS'){throw 'CURRENT_WPS_HOST_NOT_PROVEN'}
# PDF转Word由原C#宿主本地引擎提前返回；正式验收也不把它冒充WPS COM。
if($expected.feature -eq 'pdf-to-word' -and $provider -and $provider -notmatch 'WPS'){throw 'UNEXPECTED_PDF_ENGINE_PROVIDER'}
$outputRoot=[IO.Path]::GetFullPath([string]$request.output_directory).TrimEnd('\')+'\'
if(-not $outputRoot.StartsWith($root,[StringComparison]::OrdinalIgnoreCase)){throw 'NATIVE_OUTPUT_ROOT_INVALID'}
$outputs=@(foreach($path in @($jobs[0].output_paths)){
  $full=[IO.Path]::GetFullPath([string]$path)
  if(-not $full.StartsWith($outputRoot,[StringComparison]::OrdinalIgnoreCase)){throw 'NATIVE_OUTPUT_PATH_INVALID'}
  $file=Get-Item -LiteralPath $full
  if($file.Length -le 0){throw 'NATIVE_OUTPUT_EMPTY'}
  @{extension=$file.Extension.ToLowerInvariant();sha256=(Get-FileHash -LiteralPath $full -Algorithm SHA256).Hash.ToLowerInvariant();bytes=$file.Length;download_path=if($file.Extension -ieq '.pdf'){$full}else{$null}}
})
if($outputs.Count -lt 1){throw 'NATIVE_OUTPUT_REQUIRED'}
@{feature=$expected.feature;input_sha256=$expected.input_sha;host_sha256=$expected.host_sha;host_display_name=$provider;outputs=$outputs;request_sha256=(Get-FileHash -LiteralPath $requestPath -Algorithm SHA256).Hash.ToLowerInvariant();response_sha256=(Get-FileHash -LiteralPath $responsePath -Algorithm SHA256).Hash.ToLowerInvariant()}|ConvertTo-Json -Depth 5 -Compress
"""


def verify_image_derivative(pdf_path, native_sha, images, options):
    """对照产品_render_pdf_images：144DPI RGB、全页、白底居中长图、JPEG quality=92。"""
    import fitz
    from PIL import Image, ImageChops, ImageStat
    require(contract.digest(pdf_path) == native_sha, "CURRENT_NATIVE_PDF_CHANGED")
    kind, mode = options.get("target_format"), options.get("image_mode")
    require(kind in {"png", "jpg"} and mode in {"pages", "long"}
            and options.get("dpi") == 144 and options.get("page_selection") == "all", "IMAGE_DERIVATIVE_OPTIONS_MISMATCH")
    pages = []
    with fitz.open(pdf_path) as pdf:
        for page in pdf:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            pages.append(Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB"))
    require(pages, "NATIVE_PDF_PAGES_REQUIRED")
    page_sizes = [list(page.size) for page in pages]
    if mode == "long":
        canvas = Image.new("RGB", (max(page.width for page in pages), sum(page.height for page in pages)), "white")
        cursor = 0
        for page in pages:
            canvas.paste(page, ((canvas.width - page.width) // 2, cursor))
            cursor += page.height
        expected_images = [canvas]
    else:
        expected_images = pages
    require(len(images) == len(expected_images), "NATIVE_PDF_IMAGE_PAGE_COUNT_MISMATCH")
    comparisons = []
    for path, expected in zip(images, expected_images):
        if kind == "jpg":
            encoded = io.BytesIO()
            expected.save(encoded, format="JPEG", quality=92)
            encoded.seek(0)
            expected = Image.open(encoded).convert("RGB")
        with Image.open(path) as downloaded:
            require(downloaded.format == ("PNG" if kind == "png" else "JPEG"), "HTTP_IMAGE_ENCODING_MISMATCH")
            if kind == "png":
                # RGB转换会丢弃alpha；先排除透明图冒充产品的不透明PDF栅格输出。
                require(downloaded.convert("RGBA").getchannel("A").getextrema() == (255, 255), "NATIVE_PDF_PNG_ALPHA_NOT_OPAQUE")
            actual = downloaded.convert("RGB")
        require(actual.size == expected.size, "NATIVE_PDF_IMAGE_144DPI_SIZE_MISMATCH")
        diff = ImageChops.difference(expected, actual)
        mean = sum(ImageStat.Stat(diff).mean) / 3
        maximum = max(channel[1] for channel in diff.getextrema())
        histogram = diff.histogram()
        changed_ratio = sum(count for index, count in enumerate(histogram) if index % 256 > 8) / (actual.width * actual.height * 3)
        if kind == "png":
            require(diff.getbbox() is None, "NATIVE_PDF_PNG_PIXELS_DIFFER")
        else:
            # 与产品Q92重编码结果比较，容忍有限JPEG解码差异；禁止白图/错页等内容替换。
            require(mean <= 0.5 and maximum <= 32 and changed_ratio <= 0.002, "NATIVE_PDF_JPEG_PIXELS_DIFFER")
        comparisons.append({"sha256": contract.digest(path), "pixels": list(actual.size), "mean_channel_delta": mean,
                            "maximum_channel_delta": maximum, "channel_delta_over_8_ratio": changed_ratio})
    return {"status": "passed", "native_pdf_sha256": native_sha, "native_pdf_pages": len(pages),
            "page_pixels_at_144dpi": page_sizes, "dpi": 144, "mode": mode, "format": kind,
            "algorithm_source": "backend/app/official_format_features.py:_render_pdf_images",
            "algorithm_source_sha256": contract.digest(ROOT / "backend/app/official_format_features.py"),
            "png_exact": kind == "png", "jpeg_tolerance": {"mean_channel_delta_max": 0.5, "maximum_channel_delta_max": 32,
                                                    "channel_delta_over_8_ratio_max": 0.002}, "images": comparisons}


def verify_native_cases(result, proofs, host_sha):
    require(result.get("status") == "passed" and len(result.get("cases", [])) == 9 and len(proofs) == 9, "NATIVE_FEATURE_CASES_INCOMPLETE")
    cases = result["cases"]
    require([case["id"] for case in cases] == ["replace", "redheader", "rename", "convert-docx", "convert-pdf", "convert-txt", "convert-png", "convert-jpg", "pdf-to-word"]
            and len({case["job"]["id"] for case in cases}) == 9
            and {case["job"]["id"] for case in cases} == set(proofs), "CURRENT_NATIVE_CASE_SET_MISMATCH")
    seen = set()
    for case in result["cases"]:
        proof = proofs.get(case["job"]["id"], {})
        feature = case["feature_id"]
        require(proof.get("feature") == feature and proof.get("host_sha256") == host_sha
                and re.fullmatch(r"[a-f0-9]{64}", proof.get("request_sha256", ""))
                and re.fullmatch(r"[a-f0-9]{64}", proof.get("response_sha256", "")), "NATIVE_FEATURE_BINDING_MISMATCH")
        require(feature == "pdf-to-word" or "wps" in proof.get("host_display_name", "").lower(), "FEATURE_WPS_PROVIDER_REQUIRED")
        require(proof.get("outputs") and all(re.fullmatch(r"[a-f0-9]{64}", o.get("sha256", ""))
                and type(o.get("bytes")) is int and o["bytes"] > 0 for o in proof["outputs"]), "NATIVE_OUTPUT_DIGEST_REQUIRED")
        if case["id"] in {"convert-png", "convert-jpg"}:
            pdfs = [o for o in proof["outputs"] if o.get("extension") == ".pdf"]
            derivative = proof.get("image_derivative", {})
            require(len(pdfs) == 1 and derivative.get("status") == "passed"
                    and derivative.get("native_pdf_sha256") == pdfs[0]["sha256"]
                    and [o["sha256"] for o in derivative.get("images", [])] == [o["sha256"] for o in case["outputs"]],
                    "NATIVE_PDF_HTTP_IMAGE_PIXELS_NOT_BOUND")
        else:
            require({o["sha256"] for o in case["outputs"]} == {o["sha256"] for o in proof.get("outputs", [])}, "NATIVE_AND_HTTP_OUTPUT_DIFFER")
        seen.add(feature)
    require(seen == {"replace", "redheader", "rename", "convert", "pdf-to-word"}, "NATIVE_FEATURE_SET_INCOMPLETE")
    return {"status": "passed", "feature_case_count": 9, "format_native_proof": "separate-golden-jobs",
            "pdf_to_word_engine": "original-native-source-host-local-pdf-engine", "wps_features": ["format", "replace", "redheader", "rename", "convert"]}


def installation_evidence(args):
    """明确选择本次安装回执；输出仍留在原排版证据目录，不改写历史回执。"""
    selected = Path(getattr(args, "installation_evidence", None) or OUTPUT).resolve()
    allowed = {OUTPUT.resolve(), (OUTPUT.parent / "windows-timeout-repair-20261018").resolve()}
    require(selected in allowed, "INSTALLATION_EVIDENCE_DIRECTORY_NOT_ALLOWED")
    return selected


def exercise(args):
    windows = load(ROOT / "qa/vm-lab/scripts/exercise-windows-business.py", "win10_formatter_windows")
    shared = load(ROOT / "qa/vm-lab/scripts/verify-win10-x86-installed-wps.py", "win10_formatter_shared")
    # 金样/五字体共享宿主证明同样走短命令，避免新增输出断言超过cmd长度。
    shared.guest_json = lambda client, script, **kwargs: staged_readonly_probe(
        client, shared.installer.execute, script, "win10-formatter-" + uuid.uuid4().hex[:12])
    features = load(ROOT / "qa/vm-lab/guest/installed-formatter-features.py", "win10_formatter_features")
    fixtures = load(ROOT / "qa/vm-lab/scripts/exercise-linux-formatter.py", "win10_formatter_fixtures")
    transition = load(OUTPUT / "runtime-transition.py", "win10_formatter_transition")
    matrix, media = windows.load_configuration()
    matrix["defaults"].update(primary_root=str(LAB_ROOT), fallback_root=str(LAB_ROOT))
    lab = windows.QemuLab(matrix, LAB_ROOT, media)
    binding = contract.candidate(lab, ROOT, args.package, args.sha256)
    installed_evidence = installation_evidence(args)
    contract.completed_install(installed_evidence / "overlay-install", binding)
    build = contract.read(binding["package"]["build_receipt"])
    payload_manifest = Path(build["windows_payload"]["manifest"]["path"])
    require(contract.digest(payload_manifest) == binding["manifest_sha256"], "FORMAL_PAYLOAD_MANIFEST_CHANGED")
    host_entries = [r for r in contract.read(payload_manifest)["files"]
                    if r.get("path") == "formatter-host/PartyOps.DocumentFormatter.Host.exe"]
    require(len(host_entries) == 1, "FORMAL_NATIVE_HOST_BINDING_REQUIRED")
    host_sha = host_entries[0]["sha256"]
    resume_data = resume_font_evidence(args.resume_font_from, binding, host_sha, shared.oracle) if getattr(args, "resume_font_from", None) else None
    preserved = transition.preserved_inputs()[0]
    if not args.execute:
        return {"status": "offline-preflight-passed", "full_acceptance_passed": False, "guest_checked": False}
    contract.bind_guest(lab, windows.qmp)
    output = OUTPUT / ("win10-formatter-" + uuid.uuid4().hex[:12])
    output.mkdir(exist_ok=False)
    record = {"scope": "original-win10-x64-installed-http-formatter", "status": "failed", "candidate": binding,
              "generated_at": contract.now(), "full_acceptance_passed": False,
              "runtime_environment_passed": False, "six_feature_native_host_proof": "not-verified",
              "font_render_use": "not-verified", "checks": []}
    client = None
    base = prefix = headers = run = None
    observed_jobs, probes = {}, [output.name]
    record.update(stage="connect", http_failures=[], session_transport=[])
    try:
        client = windows.windows.connect(lab, matrix["targets"]["win10-x64"])
        with client.open_sftp() as fs, fs.file("C:/ProgramData/PartyOps-VM-Lab/private/" + REMOTE + ".json") as stream:
            state = json.load(stream)
        require(state.get("business_verified") is True and REMOTE in state.get("data_dir", ""), "ORIGINAL_SYNTHETIC_ACCOUNT_REQUIRED")
        run = object.__new__(windows.WindowsRun)
        run.client, run.lab, run.target, run.state = client, lab, "win10-x64", state
        run.package, run.reports = binding["package"], output
        run.args = SimpleNamespace(uuid=contract.UUID, phase="formatter")
        run.origin, run.cookies = "http://127.0.0.1:18775", http.cookiejar.CookieJar()
        # WindowsRun 原 HTTP 错误出口会写响应正文；本适配固定错误码并拒绝重定向。
        class RejectRedirect(urllib.request.HTTPRedirectHandler):
            def redirect_request(self, *_args, **_kwargs):
                raise RuntimeError("FORMATTER_REDIRECT_REJECTED")
        run.opener = urllib.request.build_opener(urllib.request.ProxyHandler({}),
            windows.GuestHTTPHandler(client.get_transport()), urllib.request.HTTPCookieProcessor(run.cookies), RejectRedirect())
        started = contract.read(installed_evidence / "overlay-install/attempt.json")["time_utc"]

        def guard(port=18775):
            contract.bind_guest(lab, windows.qmp)
            actual = contract.runtime_matches(contract.probe(client, windows.windows.execute), binding, started)
            desktop = run.desktop("Inspect")
            transition.standard_desktop(desktop)
            windows.WindowsRun.require_configuration_owner(desktop)
            row = json.loads(windows.windows.execute(client, listener_script(port)))
            validate_listener(row, actual, desktop)
            return desktop

        # 直接复用同一 tunnel opener，避免原 request 的错误正文落盘。
        def request(path, method="GET", data=None, base=None, content_type=None, extra_headers=None):
            record["last_request"] = {"route": route_template(path), "method": method if method in {"GET", "POST", "DELETE", "PUT", "PATCH"} else "UNKNOWN"}
            origin = base or run.origin
            parsed = urllib.parse.urlsplit(origin)
            require(parsed.scheme == "http" and parsed.hostname == "127.0.0.1" and parsed.port
                    and not parsed.username and not parsed.password and not parsed.path.strip("/")
                    and not parsed.query and not parsed.fragment, "GUEST_LOOPBACK_ORIGIN_REQUIRED")
            guard(parsed.port)
            require(path.startswith("/") and not path.startswith("//"), "HTTP_PATH_INVALID")
            header = {"Origin": run.origin, **(extra_headers or {})}
            for cookie in run.cookies:
                if cookie.name == "partyops_csrf": header["X-PartyOps-CSRF"] = cookie.value
            if isinstance(data, dict): data, content_type = json.dumps(data, ensure_ascii=False).encode(), "application/json"
            if content_type: header["Content-Type"] = content_type
            response_status = None
            try:
                wire_request = urllib.request.Request(origin + path, data=data, method=method, headers=header)
                with run.opener.open(wire_request, timeout=360) as response:
                    response_status = response.status
                    payload = response.read()
                    transport = session_transport_metadata(wire_request, response, payload)
                    if transport is not None:
                        record["session_transport"].append(transport)
                    value = json.loads(payload) if "application/json" in response.headers.get("Content-Type", "") else payload
                    if re.fullmatch(r"/v1/sessions/[^/]+/jobs(?:/[^/]+)?", urllib.parse.urlsplit(path).path) and isinstance(value, dict) and "id" in value:
                        try:
                            job_row = safe_job(value)
                            observed_jobs[job_row["id"]] = job_row
                        except RuntimeError:
                            pass  # 诊断不新增门禁；调用方继续执行原有断言。
                    return value
            except Exception as error:
                record["http_failures"].append({"stage": record["stage"], **http_failure(error, path, method, response_status)})
                raise RuntimeError("INSTALLED_FORMATTER_HTTP_FAILED") from None

        run.request = request
        record["stage"] = "runtime-and-wps"
        desktop = guard()
        wps = staged_readonly_probe(client, windows.windows.execute, wps_script(desktop["sid"]), output.name)
        record["checks"].append({"id": "installed-wps-x64-font-files", "result": wps})
        validate_wps(wps)
        temp = ntpath.join(ntpath.dirname(desktop["config_dir"]), "Temp")
        record["stage"] = "session"
        base, prefix, headers = shared.session(run)
        record["stage"] = "capabilities"
        capabilities = run.request("/v1/capabilities", base=base, extra_headers=headers)
        require(capabilities.get("source_host_ready") is True, "WPS_SOURCE_HOST_NOT_READY")
        import fitz
        if resume_data is not None:
            prior_checks, sources = resume_data
            record["checks"].extend(prior_checks)
            record["resume_sources"] = sources
        else:
            manifest = fixtures.feature_inputs(output)
            feature_cases = contract.read(manifest)["cases"]
            submitted, native_proofs, native_pdfs = {}, {}, {}
    
            def feature_request(origin, path, method="GET", data=None, content_type=None):
                value = run.request(path, method, data, base=origin, content_type=content_type, extra_headers=headers)
                if isinstance(value, dict) and "/jobs" in path and "/outputs/" not in path:
                    value = safe_job(value)
                    if method == "POST":
                        require(len(submitted) < len(feature_cases), "UNEXPECTED_FEATURE_JOB")
                        case = feature_cases[len(submitted)]
                        require(data["feature_id"] == case["feature_id"] and data["options"] == case["options"]
                                and len(data["document_ids"]) == 1, "FEATURE_REQUEST_CASE_MISMATCH")
                        submitted[value["id"]] = (case, data["document_ids"][0])
                    if value["state"] == "completed" and value["id"] not in native_proofs:
                        require(value["id"] in submitted, "UNBOUND_FEATURE_JOB")
                        case, doc = submitted[value["id"]]
                        script = native_host_script(temp, value["id"], doc, case["feature_id"], case["sha256"], host_sha)
                        probe_directory = "win10-formatter-" + uuid.uuid4().hex[:12]
                        probes.append(probe_directory)
                        proof = staged_readonly_probe(client, windows.windows.execute, script, probe_directory,
                                                      {"job": value["id"], "feature": case["feature_id"],
                                                       "input_sha256": case["sha256"], "host_sha256": host_sha})
                        require(proof.get("input_sha256") == case["sha256"], "FEATURE_NATIVE_INPUT_CHANGED")
                        for native_output in proof.get("outputs", []):
                            remote_pdf = native_output.pop("download_path", None)
                            if case["id"] in {"convert-png", "convert-jpg"} and native_output.get("extension") == ".pdf":
                                require(value["id"] not in native_pdfs and isinstance(remote_pdf, str)
                                        and ntpath.normpath(remote_pdf).lower().startswith(ntpath.normpath(temp).lower() + "\\")
                                        and ntpath.splitext(remote_pdf)[1].lower() == ".pdf", "CURRENT_NATIVE_PDF_PATH_INVALID")
                                local_pdf = output / ("native-" + value["id"] + ".pdf")
                                with client.open_sftp() as fs:
                                    fs.get(remote_pdf.replace("\\", "/"), str(local_pdf))
                                require(local_pdf.stat().st_size == native_output["bytes"]
                                        and contract.digest(local_pdf) == native_output["sha256"], "CURRENT_NATIVE_PDF_DOWNLOAD_CHANGED")
                                native_pdfs[value["id"]] = local_pdf
                        native_proofs[value["id"]] = proof
                return value
    
            record["stage"] = "feature-jobs-and-native-probes"
            feature_result = features.exercise(feature_request, base, prefix, manifest, output)
            record["stage"] = "feature-output-verification"
            for case in feature_result["cases"]:
                if case["id"] in {"convert-png", "convert-jpg"}:
                    job = case["job"]["id"]
                    require(job in native_pdfs, "CURRENT_NATIVE_PDF_REQUIRED")
                    case_options = next(row["options"] for row in feature_cases if row["id"] == case["id"])
                    native_sha = next(row["sha256"] for row in native_proofs[job]["outputs"] if row["extension"] == ".pdf")
                    native_proofs[job]["image_derivative"] = verify_image_derivative(native_pdfs[job], native_sha,
                        [output / row["filename"] for row in case["outputs"]], case_options)
            native_result = verify_native_cases(feature_result, native_proofs, host_sha)
            record["checks"].append({"id": "feature-native-source-host", "result": native_result, "proofs": native_proofs})
            for image in [*output.glob("convert-*.png"), *output.glob("convert-*.jpg")]:
                pixmap = fitz.Pixmap(str(image))
                require(pixmap.width > 0 and pixmap.height > 0, "CONVERT_IMAGE_DECODE_FAILED")
            record["checks"].append({"id": "six-feature-http-output-contracts", "status": feature_result["status"],
                                     "cases": feature_result["case_count"], "image_decode": "passed"})
            record["stage"] = "golden-docx"
            source = ROOT / "backend/tests/fixtures/document-formatter-source/input-manual-break.docx"
            golden = source.with_name("expected-source-formatted.docx")
            original_hashes = (contract.digest(source), contract.digest(golden))
            documents = {}
            for index in (1, 2):
                identifier, digest = shared.upload(run, base, prefix, headers, source, f"中文 空格{index}.docx")
                documents[identifier] = {"sha256": digest, "label": f"formatted-{index}"}
            outputs, proofs = shared.job(run, base, prefix, headers, "format", documents, {"scope": "full"}, output, ".docx", temp)
            oracle = shared.oracle
            signature = oracle._semantic_equivalence_signature(oracle._document_signature(golden))
            for path in outputs:
                actual = oracle._document_signature(path); oracle._require_contract(actual)
                require(oracle._semantic_equivalence_signature(actual) == signature, "GOLDEN_SEMANTIC_MISMATCH")
            record["stage"] = "golden-pdf"
            render = {}
            for index, path in enumerate((golden, *outputs)):
                identifier, digest = shared.upload(run, base, prefix, headers, path, path.name)
                render[identifier] = {"sha256": digest, "label": f"render-{index}"}
            pdfs, render_proofs = shared.job(run, base, prefix, headers, "convert", render, {"target_format": "pdf"}, output, ".pdf", temp)
            comparisons = [oracle._visual_page_comparison(pdfs[0], pdf, allow_rasterization_noise=False) for pdf in pdfs[1:]]
            require(all(row["passed"] for row in comparisons), "GOLDEN_VISUAL_MISMATCH")
            require((contract.digest(source), contract.digest(golden)) == original_hashes, "GOLDEN_INPUT_CHANGED")
            record["checks"].append({"id": "format-and-pdf-golden", "status": "passed", "host_proofs": proofs + render_proofs,
                                     "visual_comparisons": comparisons})
        record["stage"] = "font-render"
        witness = output / "five-font-witness.docx"
        font_witness(witness)
        identifier, digest = shared.upload(run, base, prefix, headers, witness, witness.name)
        witness_pdfs, witness_proofs = shared.job(run, base, prefix, headers, "convert",
            {identifier: {"sha256": digest, "label": "five-font-witness"}}, {"target_format": "pdf"}, output, ".pdf", temp)
        require(len(witness_pdfs) == 1, "FONT_WITNESS_PDF_REQUIRED")
        with fitz.open(witness_pdfs[0]) as pdf:
            traces = [span for page in pdf for span in page.get_texttrace()]
            record["font_render_observed_names"] = sorted({span.get("font", "") for span in traces})
            font_use = font_render_proof(witness_pdfs[0])
        record["checks"].append({"id": "five-font-visible-pdf-glyphs", "status": "passed", "fonts": font_use,
            "input_sha256": digest, "pdf_sha256": contract.digest(witness_pdfs[0]), "native_host_proofs": witness_proofs})
        record["stage"] = "final-runtime-verification"
        wps_after = json.loads(windows.windows.execute(client, readonly_probe_command(output.name), timeout=240))
        validate_wps(wps_after)
        require(wps_after == wps, "WPS_OR_FONT_FILES_CHANGED_DURING_RUN")
        installed_host = contract.INSTALL + r"\formatter-host\PartyOps.DocumentFormatter.Host.exe"
        final_host_sha = windows.windows.execute(client, "(Get-FileHash -LiteralPath '" + installed_host + "' -Algorithm SHA256).Hash.ToLowerInvariant()").strip()
        require(final_host_sha == host_sha, "CANDIDATE_NATIVE_HOST_CHANGED")
        record["native_host_sha256"] = final_host_sha
        guard(); require(transition.preserved_inputs()[0] == preserved, "BASELINE_CHANGED")
        record["font_render_use"] = "passed"
        record["six_feature_native_host_proof"] = "passed"
        record["status"] = "passed"
        record["formatter_acceptance_passed"] = True
        record["scope_limitation"] = "仅本次安装版HTTP排版/五字体/六入口/金样；不代替主机UAC、完整产品生命周期或其它平台验收。"
        record["executable_checks_status"] = "passed"
    except Exception as error:
        record["error"] = str(error) if isinstance(error, RuntimeError) and re.fullmatch(r"[A-Z0-9_]+", str(error)) else type(error).__name__
        record["failure_stage"] = record["stage"]
    finally:
        # Guest会话删除会清除原生job现场，必须先保存已取得的脱敏证据。
        if record["status"] != "passed":
            try: save_failure_evidence(output, record, observed_jobs, probes)
            except Exception: record["evidence_error"] = "FAILURE_EVIDENCE_WRITE_FAILED"
        record["stage"] = "cleanup"
        if base and prefix and headers:
            try: run.request(prefix, "DELETE", base=base, extra_headers=headers)
            except Exception: record.update(status="failed", cleanup_error="SESSION_CLEANUP_FAILED")
        if client is not None:
            try: client.close()
            except Exception: record.update(status="failed", cleanup_error="SSH_CLEANUP_FAILED")
        record["completed_at"] = contract.now()
        record["stage"] = record.get("failure_stage", "completed")
        record["formatter_acceptance_passed"] = record["status"] == "passed"
        (output / "summary.json").write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"status": record["status"], "full_acceptance_passed": False,
            "formatter_acceptance_passed": record["status"] == "passed", "evidence": str(output / "summary.json")}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--package", type=Path, required=True)
    parser.add_argument("--sha256", required=True)
    parser.add_argument("--execute", action="store_true", help="显式执行当前已运行安装版 HTTP，不安装或启动程序")
    parser.add_argument("--installation-evidence", type=Path, help="本次安装回执目录；仅接受已知原验收或超时修复验收目录")
    parser.add_argument("--resume-font-from", type=Path, help="只续验五字体与最终runtime；重核指定原run九功能/金样全部证据")
    try:
        result = exercise(parser.parse_args())
        print(json.dumps(result, ensure_ascii=False))
        return 0 if result["status"] in {"offline-preflight-passed", "passed"} else (3 if result["status"] == "partial" else 2)
    except Exception as error:
        print(json.dumps({"status": "failed", "error_type": type(error).__name__, "full_acceptance_passed": False}))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
